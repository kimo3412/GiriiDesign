from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(r"E:\study\GraduationProject")
OUT = ROOT / "项目文档" / "独立设计师工作室生产流程管理系统_毕业设计说明书_完整版.docx"


def set_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def fmt(p, first=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=20):
    pf = p.paragraph_format
    pf.line_spacing = Pt(line)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first:
        pf.first_line_indent = Pt(24)
    p.alignment = align


def para(doc, text, first=True, font="宋体", size=12):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, font, size)
    fmt(p, first)


def center(doc, text, font="黑体", size=18, bold=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, font, size, bold)
    fmt(p, False, WD_ALIGN_PARAGRAPH.CENTER)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    p.text = ""
    r = p.add_run(text)
    if level == 1:
        set_font(r, "黑体", 18, True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 2:
        set_font(r, "黑体", 15, True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        set_font(r, "黑体", 14, True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt(p, False, p.alignment)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def table(doc, caption, headers, rows):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_font(r, "宋体", 10.5)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_font(run, "宋体", 10.5)


def fig(doc, caption, hint):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"【此处插入图片：{hint}】")
    set_font(r, "宋体", 12)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    set_font(r, "宋体", 10.5)


def toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(end)


def add_section(doc, title, paragraphs):
    heading(doc, title, 2)
    for text in paragraphs:
        para(doc, text)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.9)
sec.bottom_margin = Inches(0.9)
sec.left_margin = Inches(1.15)
sec.right_margin = Inches(1.05)
doc.styles["Normal"].font.name = "宋体"
doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
doc.styles["Normal"].font.size = Pt(12)

center(doc, "图书分类号：TP311.5        密级：公开", "黑体", 10.5, False)
for _ in range(4):
    doc.add_paragraph()
center(doc, "独立设计师工作室生产流程管理系统的设计与实现", "黑体", 22, True)
center(doc, "DESIGN AND IMPLEMENTATION OF A PRODUCTION PROCESS MANAGEMENT SYSTEM FOR INDEPENDENT DESIGNER STUDIOS", "Arial", 16, True)
for _ in range(3):
    doc.add_paragraph()
center(doc, "毕业设计说明书", "黑体", 28, True)
for _ in range(4):
    doc.add_paragraph()
for k, v in [
    ("学生姓名", "（请填写）"),
    ("学    号", "（请填写）"),
    ("班    级", "（请填写）"),
    ("指导教师", "（请填写）"),
    ("专业名称", "软件工程"),
    ("学院名称", "信息工程学院（大数据学院）"),
]:
    center(doc, f"{k}    {v}", "黑体", 15, False)
center(doc, "2026年5月7日", "黑体", 15, False)
doc.add_page_break()

center(doc, "徐州工程学院毕业设计原创性声明", "黑体", 16, True)
para(doc, "本人郑重声明：所呈交的毕业设计，是本人在导师的指导下，独立进行研究工作所取得的成果。除文中已经注明引用或参考的内容外，本设计说明书不含任何其他个人或集体已经发表或撰写过的作品或成果。对本文的研究做出重要贡献的个人和集体，均已在文中以明确方式标注。")
para(doc, "本人完全意识到本声明的法律结果由本人承担。")
para(doc, "论文作者签名：                         日期：        年     月     日", False)
doc.add_page_break()
center(doc, "徐州工程学院毕业设计版权协议书", "黑体", 16, True)
para(doc, "本人完全了解徐州工程学院关于收集、保存、使用毕业设计的规定，即：本校学生在学习期间所完成的毕业设计的知识产权归徐州工程学院所拥有。徐州工程学院有权保留并向国家有关部门或机构送交设计说明书的纸本复印件和电子文档拷贝，允许说明书被查阅和借阅。徐州工程学院可以公布说明书的全部或部分内容，可以将说明书的全部或部分内容提交至各类数据库进行发布和检索，可以采用影印、缩印或扫描等复制手段保存和汇编本说明书。")
para(doc, "论文作者签名：                         导师签名：", False)
para(doc, "日期：        年     月     日          日期：        年     月     日", False)
doc.add_page_break()

center(doc, "摘要", "黑体", 18, True)
for text in [
    "本课题以独立设计师工作室在非标准化定制业务中的生产流程管理为研究对象，围绕服装定制、手工皮具、数字插画等小型工作室常见业务场景，分析其在客户需求采集、订单排期、生产节点记录、物料库存、客户沟通和进度反馈等方面存在的问题。传统工作室通常依赖微信群、私聊记录、Excel 表格和人工经验管理业务，信息分散且难以追溯。随着订单数量增加和业务品类扩展，传统方式容易造成需求遗漏、进度不透明、材料扣减不准确、客户反复询问和设计师协同效率低等问题。",
    "针对上述问题，本文设计并实现了一套独立设计师工作室生产流程管理系统。系统采用前后端分离架构，后端基于 Spring Boot 3.4.3、Java 17、MyBatis-Plus、MySQL、Redis、Spring Security 和 WebSocket 实现，Web 后台基于 Vue 3、TypeScript、Naive UI、Vite、Pinia 和 alova 实现，客户侧基于 Uni-app 实现微信小程序。系统由后台管理端、小程序客户端、后端服务、数据库、Redis 缓存和文件上传服务组成，形成从客户需求提交到订单完成的完整业务闭环。",
    "系统主要包括配置中心、意向管理、订单管理、节点工作台、供应链管理、作品集管理、客户管理、在线沟通、数据统计和系统管理等后台模块，以及登录、首页、作品集、定制分类、定制表单、订单列表、订单详情、聊天、个人中心、地址管理和通知中心等小程序模块。配置中心通过品类、动态字段和可配置工作流实现低代码业务扩展；意向管理实现客户需求审核和订单创建；订单管理和节点工作台实现定金支付、生产推进、节点回退、阻塞处理、尾款支付、发货和完成；供应链模块通过 BOM 模板与物料库存联动，实现订单生产过程中的材料扣减与释放；聊天模块通过 WebSocket 保存并推送客户与后台消息。",
    "系统测试从功能测试、接口测试、流程测试、异常测试和并发测试等方面展开，覆盖后台登录、菜单权限、品类字段配置、小程序提交意向、后台转单、支付定金、节点推进、阻塞解除、物料扣减、尾款支付、发货、确认收货、聊天和通知等场景。测试结果表明，系统能够满足毕业设计范围内独立设计师工作室订单全生命周期管理、低代码配置和客户进度透明化的主要需求。",
]:
    para(doc, text)
p = doc.add_paragraph()
r = p.add_run("关键词：")
set_font(r, "宋体", 12, True)
r = p.add_run("独立设计师工作室；生产流程管理；动态表单；可配置工作流；订单状态机；供应链管理；Spring Boot；Uni-app")
set_font(r, "宋体", 12)
fmt(p, False)
doc.add_page_break()

center(doc, "Abstract", "Times New Roman", 18, True)
for text in [
    "This thesis focuses on the production process management of independent designer studios in non-standard customized business scenarios, such as customized clothing, handmade leather goods and digital illustration. Many studios still rely on instant messaging records, spreadsheets and manual experience to manage customer requirements, order scheduling, production progress and material inventory. When the number of orders and business categories increases, this management mode may lead to scattered information, missing requirements, opaque progress, inaccurate material consumption and low collaboration efficiency.",
    "To solve these problems, a production process management system for independent designer studios is designed and implemented. The system adopts a front-end and back-end separated architecture. The back end is developed with Spring Boot 3.4.3, Java 17, MyBatis-Plus, MySQL, Redis, Spring Security and WebSocket. The web administration side is developed with Vue 3, TypeScript, Naive UI, Vite, Pinia and alova. The client side is implemented as a WeChat mini program based on Uni-app.",
    "The system contains modules such as configuration center, request management, order management, node workbench, supply chain management, portfolio management, customer management, online communication, statistics and system management. The mini program provides login, home page, portfolio, customization category, dynamic form, order list, order detail, chat, personal center, address management and notification center. Dynamic fields and configurable workflows enable low-code expansion. BOM templates and material inventory are linked with order production. WebSocket is used to support real-time communication between customers and designers.",
    "Functional tests, interface tests, process tests, exception tests and concurrency tests show that the system can support the complete business process from customization request submission, order conversion, deposit payment, production progress update, balance payment, delivery and receipt confirmation to order completion. The system can meet the main requirements of full life-cycle order management, low-code configuration and transparent progress tracking for independent designer studios.",
]:
    para(doc, text, font="Times New Roman")
p = doc.add_paragraph()
r = p.add_run("Key words: ")
set_font(r, "Times New Roman", 12, True)
r = p.add_run("independent designer studio; production process management; dynamic form; configurable workflow; order state machine; supply chain management; Spring Boot; Uni-app")
set_font(r, "Times New Roman", 12)
fmt(p, False)
doc.add_page_break()

center(doc, "目       录", "黑体", 18, True)
toc(doc)
para(doc, "提示：打开 Word 后右键目录区域，选择“更新域/更新整个目录”。", False)
doc.add_page_break()

heading(doc, "1 绪论", 1)
add_section(doc, "1.1 课题研究背景", [
    "独立设计师工作室是近年个性化消费和小众品牌发展背景下出现的一类轻量化生产组织。此类工作室通常由少数设计师、运营人员和手工制作人员组成，面向客户提供服装定制、手工皮具、数字插画、饰品设计等个性化服务。与标准化商品销售不同，定制业务的核心在于围绕客户需求进行设计与生产，因此业务流程具有需求差异大、沟通频次高、生产周期长、节点不固定和交付标准个性化等特点。",
    "在实际经营中，许多独立设计师工作室仍然采用较为传统的管理方式。客户通过微信、小红书、线下推荐等方式联系设计师，需求信息以聊天记录、语音、图片和临时文档形式分散保存；设计师通过 Excel 表格记录订单编号、客户姓名、报价、交付日期和当前状态；材料库存则通过人工盘点或经验估算维护。这种方式在订单数量较少时可以维持运行，但当工作室同时处理多个订单、多个品类和多个设计师任务时，问题会明显暴露。",
    "首先，客户需求缺少结构化表达。服装定制需要尺寸、面料、款式、场景和参考图，手工皮具需要皮革类型、颜色、尺寸、五金件和用途，数字插画需要画风、用途、分辨率和交付格式。如果这些信息只存在聊天记录中，后续生产阶段很容易遗漏。其次，生产进度缺少透明展示。客户通常需要反复询问“做到哪一步了”，设计师需要人工回复，增加沟通成本。再次，库存管理缺少订单级关联。材料什么时候扣减、哪个订单消耗了哪些物料、取消订单是否释放材料，都需要系统化记录。",
    "因此，构建一个面向独立设计师工作室的生产流程管理系统具有现实意义。该系统不应简单复制大型 ERP，而应更加轻量、可配置和贴近小型工作室业务。系统既要支持管理员配置品类、字段和流程，也要支持设计师处理订单和节点，还要让客户通过小程序完成需求提交、进度查看和沟通。本文围绕这一目标展开设计与实现。",
])
add_section(doc, "1.2 国内外发展现状", [
    "从企业管理信息化角度看，国内外已经形成较成熟的 ERP、MES、CRM 和 SCM 系统。ERP 侧重企业资源计划，MES 侧重生产制造执行，CRM 侧重客户关系管理，SCM 侧重供应链协同。这些系统在制造业、零售业和服务业中得到广泛应用，能够覆盖订单、生产、库存、采购、客户和财务等流程。然而，这类系统往往面向组织规模较大的企业，实施成本和学习成本较高，业务流程也偏向标准化生产。",
    "从低代码与流程平台角度看，近年来动态表单、工作流引擎和低代码平台快速发展。它们通过表单配置、流程编排和可视化管理降低系统定制成本，使企业可以在不大量编写代码的情况下调整业务流程。该思想非常适合独立设计师工作室这种品类变化频繁、字段差异明显、流程节点可配置的场景。但通用低代码平台通常缺少作品集展示、客户小程序、订单支付、物料 BOM 和生产进度透明化等行业特定功能。",
    "从移动端应用角度看，微信小程序因免安装、易传播、用户使用门槛低，在中小型服务业中应用广泛。许多预约、点餐、维修、家政和电商系统都采用小程序作为客户入口。独立设计师工作室也可以利用小程序展示作品、收集客户需求、推送订单进度和提供在线沟通。与单纯后台系统相比，小程序能够直接连接客户，使生产管理系统具备服务闭环。",
    "综合来看，现有大型管理系统功能完整但过重，通用低代码平台灵活但行业适配不足，普通电商或预约系统又缺少非标定制生产管理能力。因此，本文设计的系统将低代码配置、订单状态机、节点工作台、供应链库存、小程序客户端和 WebSocket 沟通结合起来，形成面向独立设计师工作室的专用管理系统。",
])
add_section(doc, "1.3 课题研究意义", [
    "从应用价值看，系统能够帮助独立设计师工作室建立统一的数据中心。客户资料、地址、定制意向、订单、节点进度、材料库存、聊天记录、作品集和系统配置均保存在数据库中，减少信息分散造成的查找困难。客户提交的动态表单可作为订单需求快照，设计师在生产过程中可随时查看，避免依赖聊天记录回溯。",
    "从管理价值看，系统通过品类、字段、工作流和 BOM 模板将工作室经验沉淀为可配置数据。管理员可以维护“服装定制”“手工皮具”“数字插画”等品类，为每个品类配置不同字段和生产节点。设计师按照节点工作台推进订单，使生产过程标准化但不僵化。供应链模块能够记录物料库存、预警值和 BOM 用量，降低材料管理混乱带来的风险。",
    "从客户体验看，客户无需反复询问进度，可在小程序端查看订单状态、当前节点、进度时间线和节点产出记录。支付、确认收货、地址管理和聊天均集中在小程序中完成，客户体验更加连贯。系统还能通过通知和聊天卡片提醒客户支付尾款或查看订单详情。",
    "从技术实践看，本课题覆盖软件工程中的需求分析、系统设计、数据库设计、前后端分离开发、移动端开发、接口设计、权限控制、实时通信、测试验证和文档编写等环节。系统实现过程中还考虑了重复点击、旧页面提交和库存并发扣减等实际问题，具有一定工程实践价值。",
])
add_section(doc, "1.4 课题研究内容", [
    "本文的主要研究内容包括以下几个方面。第一，分析独立设计师工作室业务流程，明确系统角色、业务边界和核心功能，建立客户、设计师、管理员三类用户模型。第二，设计系统总体架构，确定 Web 后台、小程序、后端 API、数据库、Redis 和 WebSocket 的协作方式。第三，设计动态表单与可配置工作流，使系统可以通过后台配置适配不同品类的需求字段和生产节点。",
    "第四，设计订单全生命周期流程，覆盖意向提交、后台审核、转订单、支付定金、生产推进、阻塞回退、尾款支付、发货、确认收货和取消订单。第五，设计供应链管理模块，实现物料、BOM 模板、库存扣减、库存释放和低库存预警。第六，设计聊天和通知模块，实现客户与后台实时沟通，并保存聊天历史。第七，设计系统权限模块，实现管理员、角色、菜单、字典和操作日志管理。",
])
add_section(doc, "1.5 论文组织结构", [
    "本文共分为七章。第一章为绪论，介绍课题背景、研究现状、研究意义和研究内容。第二章介绍系统开发所采用的主要技术。第三章进行系统分析，包括可行性分析、功能需求、用例分析和流程分析。第四章进行系统设计，包括总体架构、功能模块、数据库设计和核心流程设计。第五章介绍系统实现，分别说明后台、小程序、后端服务和关键业务模块。第六章介绍系统测试，包括测试环境、功能测试、异常测试、并发测试和测试结论。第七章总结全文并提出后续展望。",
])
fig(doc, "图1-1 课题技术路线图", "需求分析、系统设计、数据库设计、编码实现、测试验证路线图")

heading(doc, "2 开发技术", 1)
techs = [
    ("2.1 Java 语言", "Java 是一种面向对象的高级程序设计语言，具有跨平台、类型安全、生态成熟和工程化能力强等特点。本系统后端采用 Java 17 作为开发语言。Java 的类、接口、异常、泛型和注解机制适合构建分层清晰的业务系统，配合 Spring 生态能够快速完成 Web 接口、安全控制、数据访问和事务管理。"),
    ("2.2 Spring Boot 框架", "Spring Boot 是基于 Spring 的快速开发框架，能够通过自动配置和起步依赖减少项目搭建工作。本系统后端使用 Spring Boot 3.4.3 构建，集成 Spring Web 提供 REST API，集成 Spring Security 实现认证授权，集成 Spring WebSocket 实现在线聊天，集成 Validation 与全局异常处理提高接口健壮性。"),
    ("2.3 MyBatis-Plus", "MyBatis-Plus 是 MyBatis 的增强工具，提供通用 Mapper、ServiceImpl、条件构造器、分页插件、逻辑删除和乐观锁插件等功能。本系统中大量实体表的增删改查由 MyBatis-Plus 完成，订单、意向和物料等核心业务则结合条件更新、事务和乐观锁实现更可靠的状态变更。"),
    ("2.4 MySQL 数据库", "MySQL 是常用关系型数据库，适合保存结构化业务数据。本系统使用 MySQL 保存客户、地址、品类、字段、工作流、意向、订单、进度、物料、BOM、作品集、聊天、通知和系统权限等数据。数据库设计采用逻辑删除字段 del_flag 和乐观锁字段 version，以支持数据安全删除和并发控制。"),
    ("2.5 Redis 缓存", "Redis 是高性能内存数据库，常用于缓存、会话和令牌管理。本系统使用 Redis 保存登录 Token 和黑名单等安全状态。用户登录后，后端生成 JWT 并写入 Redis，后续请求需要携带 Token，系统可通过 Redis 校验 Token 是否有效。"),
    ("2.6 Vue 3 与 Naive UI", "Vue 3 是渐进式前端框架，支持组合式 API 和响应式数据管理。Naive UI 是 Vue 3 生态中的组件库，提供后台常用的表格、表单、弹窗、菜单、标签页、按钮和消息组件。本系统 Web 后台基于 naive-ui-admin 模板二次开发，结合 TypeScript、Pinia、Vue Router 和 alova 构建。"),
    ("2.7 Uni-app", "Uni-app 是基于 Vue 的跨端开发框架，可编译为微信小程序等平台。本系统小程序端使用 Uni-app 实现，客户可以在移动端完成登录、浏览、提交意向、查看订单、支付、确认收货和聊天等操作。小程序请求层统一封装 uni.request，自动携带 JWT Token。"),
    ("2.8 WebSocket", "WebSocket 是一种支持全双工通信的网络协议，适合实时聊天、消息推送和状态同步。本系统通过 Spring WebSocket 提供 ws://localhost:8081/ws/chat 通道，客户和后台建立连接后可以实时发送和接收消息，后端同时保存聊天历史。"),
]
for title, text in techs:
    add_section(doc, title, [text])
table(doc, "表2-1 系统开发技术及作用", ["技术", "作用", "系统应用"], [
    ["Java 17", "后端开发语言", "业务逻辑、接口服务"],
    ["Spring Boot", "后端框架", "REST API、WebSocket、配置管理"],
    ["Spring Security", "安全框架", "JWT 鉴权、权限控制"],
    ["MyBatis-Plus", "ORM 增强工具", "CRUD、分页、逻辑删除、乐观锁"],
    ["MySQL", "关系数据库", "保存核心业务数据"],
    ["Redis", "缓存数据库", "Token 和黑名单"],
    ["Vue 3", "Web 前端框架", "后台管理端"],
    ["Naive UI", "组件库", "后台表格、表单、弹窗"],
    ["Uni-app", "跨端框架", "微信小程序客户侧"],
    ["WebSocket", "实时通信协议", "客户与后台聊天"],
])

heading(doc, "3 系统分析", 1)
add_section(doc, "3.1 可行性分析", [
    "经济可行性方面，系统采用开源技术栈开发，不需要购买商业框架授权。开发环境可在普通个人电脑上搭建，部署时也可选择低成本云服务器和 MySQL、Redis 组合。对于小型工作室而言，系统能够减少人工记录和重复沟通成本，提高订单处理效率，因此具有经济可行性。",
    "技术可行性方面，系统使用的 Spring Boot、MyBatis-Plus、Vue 3、Uni-app、MySQL 和 Redis 均为成熟技术，社区资料和工程实践丰富。项目已经具备后端、Web 后台和小程序三端代码，说明技术路线能够支撑系统实现。",
    "操作可行性方面，后台采用菜单、表格、表单、弹窗、看板和工作台等常见交互，符合管理员和设计师习惯；小程序采用底部导航、卡片、表单和时间线，符合客户移动端操作习惯。复杂配置由管理员完成，客户只需按引导提交需求和查看订单。",
])
add_section(doc, "3.2 用户角色分析", [
    "系统主要包含客户、设计师和管理员三类角色。客户是定制服务购买者，主要使用小程序完成作品浏览、需求提交、订单查看、支付、确认收货和沟通。设计师是订单生产执行者，主要使用后台处理意向、查看订单、填写节点产出、推进工作流和回复客户消息。管理员是系统维护者，负责品类、字段、工作流、物料、BOM、作品集、Banner、AI 配置、角色菜单和数据统计等管理工作。",
])
table(doc, "表3-1 用户角色说明", ["角色", "使用端", "主要功能", "数据范围"], [
    ["客户", "微信小程序", "登录、作品浏览、提交意向、查看订单、支付、确认收货、聊天、地址、通知", "本人数据"],
    ["设计师", "Web 后台", "意向审核、订单处理、节点工作台、聊天、作品维护", "被分配或授权数据"],
    ["管理员", "Web 后台", "配置中心、供应链、系统管理、统计分析、客户管理", "全局管理数据"],
])
add_section(doc, "3.3 功能需求分析", [
    "系统功能可以分为客户端功能、后台业务功能、后台配置功能、供应链功能、系统管理功能和公共基础功能。客户端功能负责连接客户，后台业务功能负责订单生产流转，后台配置功能负责低代码扩展，供应链功能负责材料管理，系统管理功能负责权限和日志，公共基础功能负责登录鉴权、文件上传、通知和聊天。",
    "客户端功能包括登录、首页、作品集列表、作品详情、定制分类、定制表单、订单列表、订单详情、聊天、个人中心、资料维护、地址管理和通知中心。后台业务功能包括节点工作台、订单看板、订单列表、订单详情、意向管理和客户管理。后台配置功能包括品类、字段、工作流、Banner 和 AI 配置。供应链功能包括物料管理、BOM 模板、库存管理和出入库记录。系统管理功能包括管理员、角色、菜单、字典和操作日志。",
])
table(doc, "表3-2 后台菜单功能清单", ["一级模块", "页面", "功能说明"], [
    ["节点工作台", "/workbench/nodes", "设计师处理当前节点、保存、推进、回退、阻塞"],
    ["消息中心", "/chat/index", "客户会话、聊天历史、实时沟通"],
    ["订单管理", "/order/kanban", "订单看板、按状态或节点查看"],
    ["订单管理", "/order/list", "订单查询、详情、发货、取消等操作"],
    ["意向管理", "/request/list", "客户意向审核、转订单、关闭"],
    ["作品集", "/portfolio/list", "作品新增、编辑、发布、下架"],
    ["供应链", "/supply/material", "物料维护、库存数量、预警值"],
    ["供应链", "/supply/bom", "BOM 模板维护、品类物料用量"],
    ["供应链", "/supply/inventory", "库存查询和库存状态"],
    ["供应链", "/supply/inventory/record", "库存出入库记录"],
    ["数据分析", "/statistics/dashboard", "订单、收入、客户、库存统计"],
    ["配置中心", "/config/category", "业务品类维护"],
    ["配置中心", "/config/field", "动态字段配置"],
    ["配置中心", "/config/workflow", "工作流和节点配置"],
    ["配置中心", "/config/banner", "小程序 Banner 配置"],
    ["配置中心", "/config/ai", "AI 客服配置"],
    ["系统管理", "/system/admin", "管理员账号管理"],
    ["系统管理", "/system/role", "角色权限管理"],
    ["系统管理", "/system/menu", "菜单路由管理"],
    ["系统管理", "/system/dict", "系统字典管理"],
    ["系统管理", "/system/log", "操作日志查看"],
    ["客户管理", "/customer/list", "客户信息管理"],
    ["客户管理", "/customer/address", "客户地址查看"],
])
table(doc, "表3-3 小程序页面功能清单", ["页面", "功能说明"], [
    ["pages/login/index", "微信登录和账号测试登录"],
    ["pages/index/index", "首页 Banner、作品推荐、入口导航"],
    ["pages/portfolio/list/index", "作品集列表和品类筛选"],
    ["pages/portfolio/detail/index", "作品详情和立即定制入口"],
    ["pages/custom/category/index", "选择定制品类"],
    ["pages/custom/form/index", "动态表单填写和意向提交"],
    ["pages/order/list/index", "订单列表、状态筛选、支付、确认收货"],
    ["pages/order/detail/index", "订单详情、时间线、当前节点产出、支付"],
    ["pages/chat/index", "在线聊天、图片文件上传、转人工"],
    ["pages/user/index", "个人中心入口"],
    ["pages/user/profile/index", "头像、昵称、手机号维护"],
    ["pages/user/address/index", "地址新增、编辑、删除、默认地址"],
    ["pages/user/notification/index", "通知列表、未读数、标记已读"],
])
add_section(doc, "3.4 用例分析", [
    "系统核心用例围绕“客户提交意向—后台转单—订单生产—客户确认”展开。客户首先登录系统，浏览作品或直接选择定制品类，填写动态表单并提交意向。设计师在后台查看意向详情，根据客户需求填写报价、定金、交期和负责人，将意向转换为订单。客户支付定金后，订单进入生产中，设计师在节点工作台按工作流记录产出并推进。生产完成后，客户根据状态支付尾款，后台发货，客户确认收货。",
    "除核心业务外，管理员还需要维护基础数据，包括品类、字段、工作流、物料、BOM、作品集、Banner、角色和菜单。系统用例设计需要保证不同角色只看到自己需要的功能，客户不能访问后台数据，设计师不能越权操作不属于自己的订单，管理员则可以进行全局配置。",
])
fig(doc, "图3-1 系统用例图", "客户、设计师、管理员三类角色与系统用例")
for cap, rows in [
    ("表3-4 用例规约：提交定制意向", [["用例名称", "提交定制意向"], ["参与者", "客户"], ["前置条件", "客户已登录并选择品类"], ["基本流程", "填写动态字段、上传参考图、填写描述、提交"], ["后置条件", "生成待处理意向"], ["异常流程", "必填项为空或接口失败时提示"]]),
    ("表3-5 用例规约：审核并转订单", [["用例名称", "审核并转订单"], ["参与者", "设计师或管理员"], ["前置条件", "存在待处理意向，品类已配置工作流"], ["基本流程", "填写报价、定金、交期、设计师并提交"], ["后置条件", "创建订单并更新意向状态"], ["异常流程", "重复转单时后端拒绝"]]),
    ("表3-6 用例规约：推进节点", [["用例名称", "推进节点"], ["参与者", "设计师"], ["前置条件", "订单生产中且未阻塞"], ["基本流程", "填写节点产出，提交推进"], ["后置条件", "订单进入下一节点并新增进度"], ["异常流程", "旧页面节点快照失效时拒绝推进"]]),
]:
    table(doc, cap, ["项目", "内容"], rows)
add_section(doc, "3.5 业务流程分析", [
    "登录流程中，后台用户输入用户名和密码后，系统验证密码和账号状态，生成 JWT 并保存到 Redis。小程序客户可以通过微信登录或测试账号登录获取 Token。后续请求统一携带 Authorization 请求头，后端过滤器解析身份后放入上下文。",
    "意向转订单流程中，客户提交意向后，意向状态为待处理。后台转单时，系统查询品类对应工作流及初始节点，创建订单记录并复制客户 customData 作为需求快照。如果转单成功，意向状态改为已转单；如果重复点击转单，条件更新失败，系统不会重复创建订单。",
    "生产节点流程中，订单支付定金后进入生产中。设计师只能按照工作流 step_order 推进，不能跳过节点。保存节点表单只保存数据，不改变订单节点；推进节点会更新 current_step_id 并插入进度记录。若生产中遇到问题，设计师可阻塞订单，解除阻塞后继续推进。",
    "支付与交付流程中，订单待付定金时客户支付定金，系统扣减 BOM 物料并进入生产；生产完成后若需要尾款则进入待付尾款；客户支付尾款后进入待发货；后台发货后进入待收货；客户确认收货后进入已完成。",
])
fig(doc, "图3-2 登录流程图", "后台和小程序登录流程")
fig(doc, "图3-3 意向转订单流程图", "意向提交、审核、转单流程")
fig(doc, "图3-4 订单生产流转流程图", "支付、生产、尾款、发货、确认收货流程")

heading(doc, "4 系统设计", 1)
add_section(doc, "4.1 系统总体架构设计", [
    "系统采用前后端分离架构，整体由 Web 后台、小程序客户端、后端服务、MySQL 数据库、Redis 缓存、文件上传服务和 WebSocket 实时通信服务构成。Web 后台负责管理端业务，小程序负责客户侧业务，后端负责统一接口和业务处理，数据库负责持久化，Redis 负责认证状态，WebSocket 负责实时消息。",
    "Web 后台通过 Vite 开发服务器运行在 3100 端口，代理 /api 请求到后端 8081 端口。小程序请求层直接访问 http://localhost:8081/api。后端所有接口采用 /api/v1 前缀，后台接口多位于 /api/v1/admin 下，小程序接口多位于 /api/v1/app 下。聊天连接地址为 ws://localhost:8081/ws/chat。",
])
fig(doc, "图4-1 系统总体架构图", "Web 后台、小程序、后端服务、MySQL、Redis、文件存储、WebSocket")
add_section(doc, "4.2 后端模块设计", [
    "后端按照业务域划分为 common、system、config、order、customer、chat、portfolio、supply 和 statistics 等包。common 包提供基础设施，包括统一返回、异常处理、安全配置、JWT、Redis、WebSocket、文件上传、通知和操作日志。system 包提供 RBAC 权限管理，包括管理员、角色、菜单、字典和操作日志。config 包提供品类、字段、工作流、Banner 和 AI 配置。order 包提供意向、订单、进度和节点工作台。customer 包提供小程序用户认证、地址和通知。chat 包提供聊天历史和 WebSocket 通信。portfolio 包提供作品集。supply 包提供物料和 BOM。statistics 包提供数据统计。",
])
table(doc, "表4-1 后端包结构说明", ["包名", "主要内容", "作用"], [
    ["common", "SecurityConfig、JwtUtils、R、GlobalExceptionHandler、OssController", "基础设施和公共能力"],
    ["system", "SysAdmin、SysRole、SysMenu、Dict、OperLog", "后台权限和系统管理"],
    ["config", "DsCategory、DsCustomField、DsWorkflow、DsBanner、DsAiConfig", "业务配置中心"],
    ["order", "DsOrderRequest、DsOrder、DsOrderProgress、WorkbenchController", "意向、订单和工作台"],
    ["customer", "DsUser、DsAddress、AppAuthController、AppNotificationController", "小程序用户、地址、通知"],
    ["chat", "ChatController、ChatWebSocketHandler、DsChatMessage", "在线沟通和消息记录"],
    ["portfolio", "DsPortfolio、PortfolioController、AppPortfolioController", "作品集展示与管理"],
    ["supply", "DsMaterial、DsBomTemplate、DsBomTemplateItem", "物料和 BOM"],
    ["statistics", "StatisticsController、StatisticsService", "数据统计分析"],
])
add_section(doc, "4.3 Web 后台模块设计", [
    "Web 后台基于 naive-ui-admin 模板二次开发，采用 src/api 按业务域封装接口，src/views 存放页面，src/store 管理用户状态和动态路由。后台菜单从后端 sys_menu 表生成，登录后根据角色权限展示不同路由。页面主要包括工作台、聊天、订单、意向、作品集、供应链、统计、配置中心、客户管理和系统管理。",
    "订单相关页面包括订单看板、订单列表、订单详情和节点工作台。订单看板适合按状态或节点快速查看生产任务；订单列表适合精确查询和批量管理；订单详情展示客户信息、价格信息、工作流进度和操作按钮；节点工作台则更贴近设计师日常生产，突出当前节点表单、保存和推进操作。",
])
add_section(doc, "4.4 小程序模块设计", [
    "小程序端面向客户，采用 pages.json 配置页面和底部 TabBar。底部导航包括首页、订单和我的。首页展示 Banner、作品集和快捷入口；订单页展示客户自己的订单列表；我的页面展示个人资料、地址、通知等入口。非 TabBar 页面包括登录、作品详情、定制分类、定制表单、订单详情、聊天、地址管理、资料维护和通知中心。",
    "小程序通过 api 目录封装 auth、banner、custom、order、portfolio、user、notification、upload 等接口。utils/request.js 统一设置 BASE_URL、Authorization 请求头和错误处理。store/user.js 使用 Pinia 保存 token 和用户信息。这样的结构使页面代码主要关注交互，接口和登录状态由公共层处理。",
])
fig(doc, "图4-2 Web 后台功能结构图", "后台菜单和模块结构")
fig(doc, "图4-3 小程序功能结构图", "小程序页面和客户功能结构")
add_section(doc, "4.5 数据库总体设计", [
    "系统数据库 design_studio 包含 30 张左右的数据表，覆盖客户、地址、品类、字段、工作流、订单、进度、意向、支付、退款、账单、物料、BOM、作品集、聊天、通知、管理员、角色、菜单、字典、日志和监控等数据。业务表大多以 ds_ 开头，系统权限表以 sys_ 开头。",
    "数据库设计中，ds_category、ds_custom_field、ds_workflow 和 ds_workflow_step 组成配置模型；ds_order_request、ds_order 和 ds_order_progress 组成订单生产模型；ds_material、ds_bom_template 和 ds_bom_template_item 组成供应链模型；ds_user、ds_address 和 notification 组成客户服务模型；sys_admin、sys_role、sys_menu、sys_role_menu 和 sys_admin_role 组成权限模型。",
])
table(doc, "表4-2 数据库表清单", ["类型", "数据表", "说明"], [
    ["客户", "ds_user、ds_address", "小程序用户与收货地址"],
    ["配置", "ds_category、ds_custom_field、ds_workflow、ds_workflow_step、ds_banner、ds_ai_config", "品类、字段、流程、展示与 AI 配置"],
    ["订单", "ds_order_request、ds_order、ds_order_progress、ds_payment_record、ds_refund_record、ds_bill_record", "意向、订单、进度、支付退款和账单"],
    ["供应链", "ds_material、ds_bom_template、ds_bom_template_item、ds_bom_item", "物料和 BOM"],
    ["内容", "ds_portfolio", "作品集内容"],
    ["聊天", "ds_chat_message", "客户与后台聊天记录"],
    ["权限", "sys_admin、sys_role、sys_menu、sys_role_menu、sys_admin_role", "管理员角色菜单"],
    ["字典日志", "sys_dict_type、sys_dict_data、sys_oper_log、sys_login_log、sys_exception_log、sys_api_metrics", "字典、日志和接口监控"],
])
for caption, rows in [
    ("表4-3 ds_order 订单表主要字段", [["order_id", "订单主键"], ["order_sn", "订单编号"], ["user_id", "客户 ID"], ["category_id", "品类 ID"], ["designer_id", "设计师 ID"], ["current_step_id", "当前工作流节点"], ["status", "订单状态"], ["custom_data_snapshot", "客户需求快照"], ["total_amount", "总金额"], ["prepay_amount", "预付款"], ["paid_amount", "已付金额"], ["is_blocked", "是否阻塞"], ["version", "乐观锁版本"]]),
    ("表4-4 ds_custom_field 动态字段表主要字段", [["field_id", "字段主键"], ["category_id", "所属品类"], ["label", "显示名称"], ["field_key", "字段键"], ["field_type", "字段类型"], ["unit", "单位"], ["options", "选项 JSON"], ["placeholder", "占位提示"], ["is_required", "是否必填"], ["sort_order", "排序"]]),
    ("表4-5 ds_workflow_step 工作流节点表主要字段", [["step_id", "节点主键"], ["workflow_id", "所属工作流"], ["step_name", "节点名称"], ["step_order", "节点顺序"], ["is_initial", "是否初始节点"], ["is_final", "是否最终节点"], ["node_description", "节点说明"], ["actions", "允许动作"], ["form_schema", "节点表单"]]),
    ("表4-6 ds_material 物料表主要字段", [["material_id", "物料主键"], ["material_name", "物料名称"], ["material_code", "物料编码"], ["category", "物料类别"], ["unit", "计量单位"], ["unit_price", "单价"], ["stock_qty", "库存数量"], ["warning_qty", "预警数量"], ["version", "乐观锁版本"]]),
]:
    table(doc, caption, ["字段", "说明"], rows)
table(doc, "表4-7 主要接口设计", ["接口类别", "接口路径示例", "主要用途"], [
    ["后台认证", "/api/v1/admin/auth/login、/api/v1/admin/auth/logout、/api/v1/admin/auth/info", "后台登录、退出和获取当前用户信息。"],
    ["后台菜单", "/api/v1/admin/system/menus、/api/v1/admin/system/roles", "维护角色、菜单和权限，支撑动态路由。"],
    ["配置中心", "/api/v1/admin/config/categories、/fields、/workflows、/banners、/ai", "维护品类、字段、流程、Banner 和 AI 参数。"],
    ["意向管理", "/api/v1/admin/requests、/api/v1/admin/requests/{id}/convert", "查询客户意向、审核报价并转为正式订单。"],
    ["订单管理", "/api/v1/admin/orders、/api/v1/admin/orders/{id}", "后台查询订单列表、订单详情、取消、发货等。"],
    ["工作台", "/api/v1/admin/workbench/orders、/advance、/rollback、/block、/save", "节点工作台加载、保存节点表单、推进、回退和阻塞处理。"],
    ["供应链", "/api/v1/admin/supply/materials、/bom-templates、/inventory-records", "维护物料、BOM 模板和库存记录。"],
    ["小程序公开", "/api/v1/app/public/portfolios、/config/categories、/fields", "客户浏览作品、品类和表单配置。"],
    ["小程序业务", "/api/v1/app/requests、/orders/my、/orders/{id}/pay、/orders/{id}/confirm", "客户提交意向、查看订单、模拟支付和确认收货。"],
    ["聊天通知", "/api/v1/chat/history、ws://localhost:8081/ws/chat、/api/v1/app/notifications", "聊天历史、实时通信和通知列表。"],
])
table(doc, "表4-8 三端目录结构设计", ["端", "目录或文件", "职责"], [
    ["后端", "Code/src/main/java/com/designstudio/common", "安全、统一返回、异常、Redis、WebSocket、上传和日志等基础设施。"],
    ["后端", "Code/src/main/java/com/designstudio/order", "意向、订单、进度、工作台和订单生命周期相关接口。"],
    ["后端", "Code/src/main/java/com/designstudio/config", "品类、字段、工作流、Banner 和 AI 配置管理。"],
    ["后端", "Code/src/main/java/com/designstudio/supply", "物料、BOM 模板、库存扣减和库存记录。"],
    ["Web 后台", "web-admin/src/api", "按业务模块封装 alova 请求方法。"],
    ["Web 后台", "web-admin/src/views", "实现后台页面，包括工作台、订单、意向、供应链、配置、统计和系统管理。"],
    ["Web 后台", "web-admin/src/router 与 src/store", "处理动态路由、用户状态、标签页和权限。"],
    ["小程序", "mini-app/pages", "实现首页、作品集、定制表单、订单、聊天和个人中心页面。"],
    ["小程序", "mini-app/api 与 mini-app/utils/request.js", "封装小程序接口请求、Token 携带和错误提示。"],
    ["小程序", "mini-app/store/user.js", "保存客户登录状态和用户资料。"],
])
fig(doc, "图4-4 数据库 E-R 图", "用户、意向、订单、流程、物料、聊天和权限核心 E-R 图")
add_section(doc, "4.6 订单状态机设计", [
    "订单状态机是系统的核心。订单状态 status 用于描述订单宏观阶段，current_step_id 用于描述生产中的具体节点。这样既能让客户看到简单的订单阶段，又能让设计师在后台管理详细生产步骤。订单状态包括待付定金、生产中、待发货、待收货、已完成、已取消和待付尾款。",
    "状态流转必须由后端统一控制。客户只能在待付定金和待付尾款状态执行支付，在待收货状态执行确认收货；设计师只能在生产中通过工作台推进节点；管理员或后台只能在符合条件时发货、取消或延期。通过这种方式，系统避免前端直接修改状态造成混乱。",
])
table(doc, "表4-9 订单状态说明", ["状态值", "状态名称", "说明"], [
    ["0", "待付定金", "订单已创建，等待客户支付定金"],
    ["1", "生产中", "订单进入工作流节点生产"],
    ["2", "待发货", "生产完成且款项满足发货条件"],
    ["3", "待收货", "后台已发货，等待客户确认"],
    ["4", "已完成", "客户确认收货后完成"],
    ["5", "已取消", "订单已取消"],
    ["6", "待付尾款", "生产完成后等待客户支付尾款"],
])
fig(doc, "图4-5 订单状态转换图", "订单状态机转换图")

heading(doc, "5 系统实现", 1)
add_section(doc, "5.0 模块协作实现概述", [
    "系统不是由单个订单模块孤立完成业务，而是由配置、客户、意向、订单、供应链、聊天、通知和权限等模块协作完成。配置中心先建立品类、字段和工作流，供应链模块为品类维护 BOM 模板，作品集模块为客户提供展示入口。客户在小程序中选择品类并填写动态表单后，意向模块保存需求。后台审核意向时，订单模块读取品类对应工作流并生成订单。客户支付定金后，订单模块调用物料扣减逻辑，使供应链库存与生产订单关联起来。设计师推进节点时，工作台保存节点表单和进度记录，小程序订单详情通过时间线展示给客户。聊天和通知模块贯穿订单全程，用于提醒客户付款、确认收货或补充需求。",
    "这种模块协作方式体现了本系统的业务闭环：配置决定可售卖的定制服务，客户入口产生需求，后台转化为订单，订单驱动生产节点，生产节点驱动物料和进度，进度再反馈给客户。系统管理模块位于最底层，负责账号、角色、菜单和日志，保证不同角色在各自边界内使用系统。"
])
impl_sections = [
    ("5.1 后端公共基础实现", [
        "公共基础模块是系统稳定运行的基础。统一返回对象 R<T> 封装 code、msg 和 data，使前后端对接口结果有统一判断标准。GlobalExceptionHandler 统一处理业务异常和运行时异常，避免异常直接暴露到前端。BusinessException 用于主动抛出业务错误，例如订单状态不允许操作、库存不足或节点已变化。",
        "安全配置由 SecurityConfig、JwtAuthenticationFilter、JwtUtils、LoginHelper 和 LoginUser 组成。用户登录后获得 JWT，后续请求由过滤器解析 Token 并将用户信息写入上下文。LoginHelper 提供获取当前用户 ID 和用户类型的方法，业务代码无需重复解析 Token。RedisConfig 提供 Redis 序列化和连接配置，用于 Token 管理和黑名单。",
        "文件上传由 OssController 提供，支持普通上传和聊天文件上传。StaticResourceConfig 配置本地上传目录映射，使前端能够访问上传后的图片或文件。OperLog 注解和 OperLogAspect 用于记录后台关键操作，SysOperLogController 提供日志查询。",
    ]),
    ("5.2 系统权限模块实现", [
        "系统权限模块采用 RBAC 模型。sys_admin 保存管理员账号，sys_role 保存角色，sys_menu 保存菜单和权限标识，sys_admin_role 表示管理员与角色关系，sys_role_menu 表示角色与菜单关系。后台登录后，根据用户角色查询菜单树，前端再根据菜单动态生成路由。",
        "管理员管理页面支持账号新增、编辑、禁用和角色分配。角色管理页面支持角色新增、编辑和菜单授权。菜单管理页面维护路由路径、组件路径、权限标识、排序和显示状态。字典模块维护系统中常用的字典类型和字典项，例如状态值、分类值等。操作日志模块记录关键管理行为，便于后期审计。",
    ]),
    ("5.3 配置中心实现", [
        "配置中心是系统低代码能力的核心。品类管理维护 ds_category，字段配置维护 ds_custom_field，工作流配置维护 ds_workflow 和 ds_workflow_step。管理员可以新增“服装定制”“手工皮具”“数字插画”等品类，并为每个品类配置不同字段和工作流。",
        "字段配置支持 text、number、textarea、date、select、radio、checkbox 和 image 等类型。字段具有 label、fieldKey、unit、placeholder、options、isRequired 和 sortOrder 等属性。小程序定制表单页根据字段类型动态渲染控件，提交时将用户填写的数据序列化为 JSON。",
        "工作流配置支持为每个品类设置多个生产节点，例如服装定制可包含需求确认、面料采购、裁剪、缝制、质检等节点；皮具可包含需求确认、皮料准备、裁片、缝制、封边、包装等节点；数字插画可包含需求确认、草稿、线稿、上色、交付等节点。节点的 step_order 决定推进顺序。",
    ]),
    ("5.4 小程序客户服务实现", [
        "小程序登录页提供微信登录和账号测试登录。微信登录通过 uni.login 获取 code 并调用 /v1/app/auth/wx-login，测试登录调用 /v1/app/auth/mock-login。登录成功后，用户 Token 和用户信息保存到 Pinia store 和本地缓存中。请求封装层在每次请求时自动读取 token 并放入 Authorization 请求头。",
        "首页 pages/index/index.vue 加载 Banner 和作品集数据，为客户提供作品浏览和定制入口。作品集列表和详情调用 /v1/app/public/portfolios 接口，只展示后台已发布作品。客户可以从作品详情进入定制流程，也可以从首页直接选择定制分类。",
        "个人中心提供资料、地址和通知入口。资料页支持头像上传、昵称和手机号维护；地址页支持新增、编辑、删除和默认地址；通知中心支持通知列表、未读数量、单条已读和全部已读。上述接口均从登录上下文获取 userId，保证客户只能操作自己的数据。",
    ]),
    ("5.5 定制意向模块实现", [
        "客户进入定制分类页后，小程序调用 /v1/app/public/config/categories 获取启用品类。选择品类后进入动态表单页，页面调用 /v1/app/public/config/categories/{categoryId}/fields 获取字段 Schema。页面根据字段类型渲染输入框、数字输入、文本域、日期选择、单选、多选、下拉选择和图片控件。",
        "提交意向时，小程序将 categoryId、customData、imageUrls 和 description 发送到 /v1/app/requests。AppRequestController 从 LoginHelper 获取当前客户 ID，调用 RequestService 保存意向。保存后的意向状态为待处理，后台意向管理页面可看到该记录。",
        "后台意向列表支持按状态、品类和客户查询。设计师审核意向时，需要填写总金额、预付款、预计交期、设计师和备注等信息。RequestServiceImpl 在转单时创建订单，复制客户需求快照，设置初始状态和工作流节点，并将意向状态更新为已转单。为防止重复转单，服务层会检查更新影响行数。",
    ]),
    ("5.6 订单管理模块实现", [
        "订单管理包括订单列表、订单看板和订单详情。订单列表适合查询和管理全部订单，支持状态、品类、客户和设计师等条件筛选。订单看板适合按状态或节点查看生产任务，帮助设计师快速掌握当前工作量。订单详情展示订单基础信息、客户需求、价格、当前状态、当前节点、进度记录和操作按钮。",
        "小程序订单列表调用 /v1/app/orders/my，只返回当前客户的订单。列表按状态显示进行中、待收货、已完成和全部订单。订单详情调用 /v1/app/orders/{id}/timeline，展示订单状态、工作流节点、当前节点产出、进度记录、阻塞原因、延期原因和支付按钮。",
        "支付接口 /v1/app/orders/{id}/pay 根据订单状态判断支付定金或尾款。待付定金时，支付成功后订单进入生产中并扣减 BOM 物料；待付尾款时，支付成功后订单进入待发货。确认收货接口 /v1/app/orders/{id}/confirm 只允许订单处于待收货状态时执行。",
    ]),
    ("5.7 节点工作台实现", [
        "节点工作台是设计师处理订单生产的核心页面。设计师可以按品类筛选生产中订单，选择订单后查看当前节点、客户需求、节点表单和历史产出。工作台支持 save、advance、rollback、block、unblock 等动作。",
        "保存动作只保存当前节点表单数据，不推进节点。推进动作会校验订单处于生产中、未阻塞、目标节点是当前节点的合法后继，并检查前端传入的 expectedCurrentStepId 是否与数据库 current_step_id 一致。如果订单已经被其他页面推进，旧页面提交会被拒绝。",
        "回退动作用于处理返工场景，会将订单退回到上一个节点并新增进度记录。阻塞动作用于标记生产中遇到的问题，例如客户未确认、材料不足或沟通暂停。解除阻塞后订单可以继续推进。通过这些动作，系统能够覆盖实际生产中常见的正常推进和异常处理。",
    ]),
    ("5.8 供应链管理实现", [
        "供应链模块包括物料管理、BOM 模板、库存管理和库存记录。物料表保存物料名称、编码、类别、单位、单价、库存数量和预警数量。管理员可以维护面料、皮料、五金、辅料和包装材料。低库存查询用于发现即将不足的物料。",
        "BOM 模板用于描述某个品类生产一件订单通常需要消耗的物料及数量。例如服装定制可能需要面料、衬里、拉链、暗扣和包装盒；手工皮具可能需要植鞣革、蜡线、五金扣和包装盒。订单进入生产时，系统根据品类查询 BOM 模板并扣减物料库存。",
        "物料库存更新使用乐观锁保护。若两个订单同时扣减同一物料，只有 version 匹配的更新会成功；失败时服务层抛出异常，避免页面提示成功但库存实际未变化。订单取消时，系统按订单已分配材料释放库存，保证库存数据与订单状态保持一致。",
    ]),
    ("5.9 聊天与通知实现", [
        "聊天模块通过 ChatWebSocketHandler 实现实时通信。客户进入聊天页后连接 ws://localhost:8081/ws/chat?token=...，后端校验 Token 后将会话加入 SessionManager。客户发送消息时，消息体包含 type、content、msgType、extraJson 和 orderId。后端保存 DsChatMessage 后将消息推送给后台管理员。",
        "后台消息中心显示客户会话列表和聊天历史，管理员可回复客户。聊天支持文本、图片和文件，文件通过 /v1/oss/chat-upload 上传。系统还支持转人工请求和 AI 辅助回复，当客户询问订单、进度、付款等关键词时，后端可以推送订单进度卡片或操作卡片。",
        "通知模块用于保存客户业务通知，例如订单节点更新、付款提醒和交付提醒。小程序通知中心可以获取通知列表、未读数量，支持单条标记已读和全部标记已读。通知与聊天共同构成客户沟通闭环。",
    ]),
    ("5.10 作品集与内容展示实现", [
        "作品集模块用于展示工作室作品和案例。后台作品集页面支持新增、编辑、上传图片、设置品类、排序、发布和下架。只有发布状态的作品会通过 /v1/app/public/portfolios 接口展示给小程序客户。作品详情展示标题、描述、图片和浏览量。",
        "Banner 模块用于管理小程序首页轮播图。管理员可以上传 Banner 图片、设置跳转链接、排序和状态。小程序首页调用 /v1/app/public/banners 获取启用的 Banner。作品集和 Banner 共同承担客户获客与内容展示功能。",
    ]),
    ("5.11 数据统计实现", [
        "数据统计模块聚合订单、客户、收入和库存等数据。后台统计页面展示订单数量、进行中订单、已完成订单、收入金额、客户数量、品类分布、订单趋势和库存预警。统计数据来自订单表、客户表和物料表，通过 StatisticsServiceImpl 聚合后返回给前端。",
        "对独立设计师工作室而言，统计模块能够帮助管理员了解当前业务情况，例如哪些品类订单较多、哪些订单逾期、哪些材料库存不足、近期收入变化如何等。虽然本毕业设计中的统计功能仍较基础，但已经具备经营看板的雏形。",
    ]),
    ("5.12 Web 后台页面实现", [
        "Web 后台基于 Vue 3 组合式 API 编写，页面大量使用 Naive UI 的 n-data-table、n-form、n-modal、n-drawer、n-select、n-upload 和 n-tabs 等组件。列表页通常由查询表单、工具栏、数据表格和分页组成；编辑页通常使用弹窗或抽屉承载表单；详情页则以描述列表、时间线和操作按钮展示核心信息。这样的页面结构符合后台管理系统高频查询和批量维护的使用习惯。",
        "接口请求统一放在 src/api 目录，页面不直接拼接 URL，而是调用封装好的方法。请求层基于 alova，结合统一响应拦截处理登录失效、业务错误和成功提示。用户登录后，Pinia user store 保存 Token、用户信息和权限；路由模块根据后端菜单生成动态路由，并通过 permission 指令控制按钮级权限。"
    ]),
    ("5.13 小程序页面实现", [
        "小程序端以客户体验为中心，页面层级尽量简洁。首页负责展示内容和入口，定制表单负责收集需求，订单详情负责展示进度和操作，聊天页负责沟通，个人中心负责账号、地址和通知。页面通过 pages.json 注册，并使用底部 TabBar 连接首页、订单和我的三个高频入口。",
        "小程序请求封装层处理 BASE_URL、Token 注入、加载提示和错误提示，页面只关注业务数据。动态表单页根据字段类型切换不同控件，图片字段通过上传接口返回 URL 后写入表单数据；订单详情页根据 status 控制按钮显示，例如待付定金显示支付定金，待付尾款显示支付尾款，待收货显示确认收货。"
    ]),
    ("5.14 数据一致性实现", [
        "订单系统最容易出现的问题是重复提交、旧页面提交和库存并发扣减。为此，系统在关键写操作中加入状态校验和乐观锁思想。意向转订单时，只有待处理意向允许转单；节点推进时，前端提交 expectedCurrentStepId，后端与数据库 current_step_id 比对；物料扣减时，更新条件包含 version 和库存数量。",
        "这些处理并不是为了让系统变复杂，而是为了保证页面点击、网络重试或多窗口操作不会破坏订单状态。对于毕业设计系统来说，这部分实现能够体现工程化思维：前端交互可以友好，但真正的业务边界必须由后端控制，数据库更新结果也必须被检查。"
    ]),
]
for title, paragraphs in impl_sections:
    add_section(doc, title, paragraphs)
for caption, hint in [
    ("图5-1 小程序首页", "首页 Banner、作品集和定制入口截图"),
    ("图5-2 定制分类页面", "客户选择服装、皮具、插画等品类截图"),
    ("图5-3 动态定制表单页面", "根据字段配置渲染的表单截图"),
    ("图5-4 小程序订单列表页面", "客户订单列表和支付按钮截图"),
    ("图5-5 小程序订单详情页面", "订单时间线、节点产出和支付按钮截图"),
    ("图5-6 后台意向管理页面", "意向列表和审核转单弹窗截图"),
    ("图5-7 后台订单看板页面", "订单看板截图"),
    ("图5-8 后台订单详情页面", "订单详情和进度记录截图"),
    ("图5-9 节点工作台页面", "设计师节点表单、保存和推进截图"),
    ("图5-10 品类与字段配置页面", "配置中心品类和字段页面截图"),
    ("图5-11 工作流配置页面", "工作流节点配置截图"),
    ("图5-12 物料管理页面", "物料库存列表截图"),
    ("图5-13 BOM 模板页面", "品类 BOM 物料用量截图"),
    ("图5-14 作品集管理页面", "后台作品集列表截图"),
    ("图5-15 在线聊天页面", "小程序或后台聊天截图"),
    ("图5-16 数据统计页面", "统计看板截图"),
    ("图5-17 系统角色菜单页面", "角色授权或菜单管理截图"),
]:
    fig(doc, caption, hint)

heading(doc, "6 系统测试", 1)
add_section(doc, "6.1 测试环境", [
    "系统测试在本地开发环境中进行，后端服务运行在 8081 端口，Web 后台运行在 3100 端口，小程序编译目标为微信小程序。测试前启动 MySQL 和 Redis，并导入初始化脚本和测试数据。后台测试账号为 admin/admin123，小程序可使用测试登录入口。",
])
table(doc, "表6-1 测试环境", ["项目", "配置"], [
    ["操作系统", "Windows 10/11"],
    ["JDK", "Java 17"],
    ["后端框架", "Spring Boot 3.4.3"],
    ["数据库", "MySQL 8.0"],
    ["缓存", "Redis"],
    ["后台端口", "8081"],
    ["Web 前端端口", "3100"],
    ["小程序目标", "mp-weixin"],
])
add_section(doc, "6.2 功能测试", [
    "功能测试按照系统真实业务流程展开，既测试单个页面功能，也测试跨模块流程。重点验证后台菜单是否可进入、列表是否能加载、表单是否能保存、状态是否按预期变化、小程序是否能访问客户自己的数据、核心写操作是否能正确提示成功或失败。",
])
table(doc, "表6-2 功能测试用例", ["编号", "测试项", "操作步骤", "预期结果"], [
    ["TC01", "后台登录", "输入 admin/admin123 登录", "登录成功，菜单加载"],
    ["TC02", "品类配置", "新增或编辑测试品类", "品类保存成功，小程序可读取启用品类"],
    ["TC03", "字段配置", "给品类新增动态字段", "小程序表单显示对应字段"],
    ["TC04", "工作流配置", "维护品类工作流节点", "订单生产节点按顺序显示"],
    ["TC05", "提交意向", "小程序填写表单并提交", "后台意向列表出现记录"],
    ["TC06", "转订单", "后台审核意向并转单", "创建订单，意向状态已转单"],
    ["TC07", "支付定金", "小程序支付定金", "订单进入生产中并扣减物料"],
    ["TC08", "节点保存", "工作台填写表单后保存", "保存产出，节点不变化"],
    ["TC09", "节点推进", "工作台点击推进", "进入下一节点并新增进度"],
    ["TC10", "节点回退", "执行回退操作", "回到上一节点并记录返工"],
    ["TC11", "阻塞解除", "阻塞订单后解除阻塞", "阻塞状态正确变化"],
    ["TC12", "尾款支付", "完工后客户支付尾款", "订单进入待发货"],
    ["TC13", "发货确认", "后台发货，客户确认收货", "订单进入已完成"],
    ["TC14", "聊天", "客户发送消息，后台回复", "双方消息实时显示并落库"],
    ["TC15", "地址管理", "新增、编辑、删除地址", "地址列表正确更新"],
    ["TC16", "作品集", "发布作品后小程序查看", "已发布作品显示"],
])
add_section(doc, "6.3 异常测试", [
    "异常测试主要验证系统对错误状态和非法操作的处理能力。系统应在前端给出明确提示，在后端保持数据一致，不能出现业务操作失败但页面提示成功的情况。尤其是订单推进、意向转单和物料扣减等核心写操作，需要检查数据库更新结果。",
])
table(doc, "表6-3 异常测试用例", ["编号", "异常场景", "预期结果"], [
    ["E01", "未登录访问我的订单", "返回未登录并跳转登录页"],
    ["E02", "动态表单必填项为空", "小程序提示补充必填字段"],
    ["E03", "重复点击意向转单", "只生成一个订单，重复操作失败"],
    ["E04", "旧页面推进订单节点", "后端拒绝并提示状态已变化"],
    ["E05", "库存不足时支付定金", "订单不进入生产，提示库存不足"],
    ["E06", "已完成订单再次支付", "提示当前状态不允许支付"],
    ["E07", "阻塞订单继续推进", "提示先解除阻塞"],
    ["E08", "无权查看他人订单", "接口返回订单不存在或无权查看"],
])
add_section(doc, "6.4 并发测试", [
    "并发测试围绕重复点击和两页面同时操作展开。同一订单在两个页面打开，页面 A 推进节点成功后，页面 B 不刷新继续推进。由于页面 B 携带 expectedCurrentStepId 与数据库 current_step_id 不一致，后端拒绝推进，避免订单一次操作跳过两个节点。同一意向连续点击转单时，后端条件更新保证只有第一次成功。同一物料被多个订单扣减时，version 乐观锁保证库存更新一致。",
])
add_section(doc, "6.5 构建测试", [
    "后端执行 .\\mvnw.cmd clean test 验证编译和测试；Web 后台执行 pnpm run build 验证生产构建；小程序执行 npm run build:mp-weixin 验证小程序编译。构建测试能够发现语法错误、类型错误、依赖缺失和打包配置问题。",
])
add_section(doc, "6.6 测试总结", [
    "测试结果表明，系统能够完成毕业设计范围内的主要业务闭环。后台配置、小程序提交、意向转单、订单支付、节点推进、物料扣减、聊天沟通和统计展示等功能能够协同运行。系统仍有可完善之处，例如真实支付、物流接口、图片长期存储、自动化测试覆盖率和性能压力测试等，但不影响当前核心功能展示。",
])

heading(doc, "结论", 1)
para(doc, "本文针对独立设计师工作室非标准化定制业务中的需求分散、流程不透明、客户沟通成本高、订单排期困难和物料库存管理粗放等问题，设计并实现了一套生产流程管理系统。系统基于 Spring Boot、MyBatis-Plus、MySQL、Redis、Vue 3、Naive UI 和 Uni-app 构建，覆盖后台管理端、小程序客户端和后端服务。")
para(doc, "系统通过配置中心实现品类、动态字段、工作流、Banner 和 AI 配置管理，通过意向管理和订单管理实现客户需求到正式订单的流转，通过节点工作台实现生产过程记录和节点推进，通过供应链模块实现物料与 BOM 管理，通过聊天和通知模块实现客户与后台沟通，通过系统管理模块实现 RBAC 权限控制。整体上，系统形成了从客户提交定制意向到订单完成的完整闭环。")
para(doc, "本系统的创新点主要体现在三个方面。第一，将动态表单和可配置工作流结合到独立设计师工作室场景中，降低新增业务品类的开发成本。第二，将订单状态机与生产节点结合，使客户侧和设计师侧都能获得合适粒度的进度信息。第三，在订单推进、意向转单和物料扣减等关键操作中加入乐观锁、状态快照和更新结果检查，提高系统可靠性。")
para(doc, "由于时间和条件限制，系统仍存在不足。例如支付功能目前为模拟支付，尚未接入真实微信支付；物流发货未对接第三方物流平台；图片上传和长期存储策略仍可优化；工作流配置还可以进一步做成可视化拖拽；自动化测试和压力测试覆盖面有待提升。后续可以从真实部署、性能优化、智能排期、移动端体验和数据分析深度等方向继续完善。")

heading(doc, "参考文献", 1)
for ref in [
    "[1] Craig Walls. Spring in Action[M]. 6th ed. Manning Publications, 2022.",
    "[2] MyBatis-Plus 官方文档[EB/OL]. https://baomidou.com/.",
    "[3] Spring Boot Reference Documentation[EB/OL]. https://docs.spring.io/spring-boot/.",
    "[4] Vue.js 官方文档[EB/OL]. https://cn.vuejs.org/.",
    "[5] Naive UI 官方文档[EB/OL]. https://www.naiveui.com/.",
    "[6] DCloud. Uni-app 官方文档[EB/OL]. https://uniapp.dcloud.net.cn/.",
    "[7] Redis Documentation[EB/OL]. https://redis.io/docs/.",
    "[8] MySQL 8.0 Reference Manual[EB/OL]. https://dev.mysql.com/doc/.",
    "[9] Martin Fowler. Patterns of Enterprise Application Architecture[M]. Addison-Wesley, 2002.",
    "[10] Gamma E, Helm R, Johnson R, Vlissides J. Design Patterns[M]. Addison-Wesley, 1994.",
    "[11] 阮一峰. RESTful API 设计指南[EB/OL]. http://www.ruanyifeng.com/blog/2014/05/restful_api.html.",
    "[12] 李刚. 轻量级 Java EE 企业应用实战[M]. 北京：电子工业出版社, 2021.",
    "[13] 王珊, 萨师煊. 数据库系统概论[M]. 北京：高等教育出版社, 2014.",
    "[14] 周志明. 深入理解 Java 虚拟机[M]. 北京：机械工业出版社, 2019.",
    "[15] Eric Evans. Domain-Driven Design[M]. Addison-Wesley, 2003.",
]:
    para(doc, ref, False)
heading(doc, "致谢", 1)
para(doc, "本毕业设计从选题、需求分析、系统设计、编码实现到论文撰写，得到了指导教师的耐心指导和帮助。老师在课题定位、业务流程梳理、系统模块划分和论文结构安排等方面提出了许多宝贵意见，使本人能够更加清晰地理解软件工程项目从需求到实现再到测试的完整过程。")
para(doc, "感谢学院提供的软件工程课程学习环境，使本人能够将 Java Web 开发、数据库设计、前端开发、移动端开发、软件测试和项目管理等知识综合应用于本课题。感谢同学在系统功能体验、测试数据准备和文档排版方面给予的建议和帮助。通过本次毕业设计，本人不仅完成了一个较为完整的前后端分离系统，也对业务建模、工作流设计、状态机控制和工程化开发有了更深入的理解。")
heading(doc, "附录", 1)
add_section(doc, "附录1 系统运行命令", [
    "后端运行：cd Code && .\\mvnw.cmd spring-boot:run",
    "后端测试：cd Code && .\\mvnw.cmd clean test",
    "Web 后台运行：cd web-admin && pnpm run dev",
    "Web 后台构建：cd web-admin && pnpm run build",
    "小程序运行：cd mini-app && npm run dev:mp-weixin",
    "小程序构建：cd mini-app && npm run build:mp-weixin",
])
add_section(doc, "附录2 图片清单", [
    "本文预留图片包括技术路线图、用例图、登录流程图、意向转订单流程图、订单生产流程图、总体架构图、后台功能结构图、小程序功能结构图、数据库 E-R 图、订单状态转换图、小程序首页、定制表单、订单列表、订单详情、后台意向管理、订单看板、节点工作台、物料与 BOM、作品集、聊天、统计和角色菜单等。用户后续可根据实际系统截图替换占位内容。",
])

for section in doc.sections:
    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.text = ""
    run = foot.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
